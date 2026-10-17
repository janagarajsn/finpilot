"""Import pipeline against a real DB: template round-trips, preview, duplicates, FX, undo."""

from __future__ import annotations

import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from finpilot.categories import service as cats
from finpilot.db.session import user_session
from finpilot.ingest import import_service as imp
from finpilot.ingest import template
from finpilot.money import get_currency_registry
from tests.conftest import needs_pg, needs_redis
from tests.helpers import StubFX, make_user

pytestmark = [needs_pg, needs_redis]


@pytest.fixture
def uid(redis_client):
    return make_user("INR").user_id


def upload(uid, content: bytes, name="t.csv"):
    with user_session(uid) as s:
        return imp.create_batch(s, uid, name, content)["id"]


def do_import(uid, bid, fx=None):
    fx = fx or StubFX()
    with user_session(uid) as s:
        imp.mark_importing(s, uid, bid)
    with user_session(uid) as s:
        return imp.run_import(s, uid, bid, fx)


def count(uid) -> int:
    with user_session(uid) as s:
        return s.execute(text("SELECT count(*) FROM transactions")).scalar_one()


def test_csv_template_imports_with_zero_errors_and_reupload_is_duplicate(uid):
    content = template.build_csv("INR")
    bid = upload(uid, content, "template.csv")
    with user_session(uid) as s:
        pv = imp.preview(s, uid, bid, StubFX())
    assert pv["counts"]["invalid"] == 0 and pv["counts"]["new"] == 4, pv["errors"]
    usd = next(r for r in pv["rows"] if r["currency"] == "USD")
    assert usd["amount_inr"] == Decimal("-1663.00") and usd["fx_source"] == "mcp"  # 20 × 83.15
    counts = do_import(uid, bid)
    assert counts["imported"] == 4 and count(uid) == 4
    with user_session(uid) as s:
        cats_ = {
            r[0]
            for r in s.execute(
                text("SELECT c.name FROM transactions t JOIN categories c ON c.id = t.category_id")
            )
        }
        assert {"Salary", "Rent/Mortgage", "Subscriptions"} <= cats_
        assert (
            s.execute(text("SELECT count(*) FROM import_files")).scalar_one() == 0
        )  # raw upload deleted
    # the same file again: everything is a duplicate and nothing is imported
    bid2 = upload(uid, content, "template.csv")
    with user_session(uid) as s:
        assert imp.preview(s, uid, bid2, StubFX())["counts"]["duplicate"] == 4
    assert do_import(uid, bid2)["imported"] == 0 and count(uid) == 4


def test_xlsx_template_roundtrip_has_sheets_dropdowns_and_imports(uid):
    with user_session(uid) as s:
        categories = cats.list_categories(s, uid)
    data = template.build_xlsx(categories, get_currency_registry().codes, "USD")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames[:3] == ["Transactions", "Instructions", "Categories"]
    dvs = {str(dv.formula1) for dv in wb["Transactions"].data_validations.dataValidation}
    assert (
        '"expense,income"' in dvs
        and any("Categories!" in f for f in dvs)
        and any("Lists!" in f for f in dvs)
    )
    bid = upload(uid, data, "template.xlsx")
    with user_session(uid) as s:
        pv = imp.preview(s, uid, bid, StubFX())
    assert pv["counts"]["invalid"] == 0 and pv["counts"]["new"] == 4, pv["errors"]
    assert do_import(uid, bid)["imported"] == 4


def test_ambiguous_dates_require_choice_and_bank_statement_mapping(uid):
    content = (
        b"Txn Date,Narration,Withdrawal Amt,Deposit Amt\n"
        b"05/01/2024,UPI/P2M/1234567890/SWIGGY/x@ybl,450.00,\n"
        b"06/01/2024,NEFT-HDFC0001234-ACME LTD-SALARY,,85000.00\n"
        b"06/01/2024,UPI/P2M/1234567890/SWIGGY/x@ybl,450.00,\n"
    )
    bid = upload(uid, content)
    with user_session(uid) as s:
        assert imp.preview(s, uid, bid, StubFX())["needs_date_format"] is True
        imp.update_mapping(s, uid, bid, None, {"date_format": "%d/%m/%Y"})
        pv = imp.preview(s, uid, bid, StubFX())
    assert pv["counts"] == {
        "total": 3,
        "valid": 3,
        "new": 3,
        "duplicate": 0,
        "invalid": 0,
        "warnings": 0,
    }
    assert pv["rows"][0]["date"].isoformat() == "2024-01-05" and pv["rows"][1]["kind"] == "income"
    do_import(uid, bid)
    with user_session(uid) as s:
        merchants = {r[0] for r in s.execute(text("SELECT merchant_norm FROM transactions"))}
    assert merchants == {"SWIGGY", "ACME"}


def test_identical_rows_in_one_file_both_import_and_preview_flags_existing(uid):
    rows = "date,description,amount,type\n2024-02-01,Chai,20,expense\n2024-02-01,Chai,20,expense\n"
    bid = upload(uid, rows.encode())
    assert do_import(uid, bid)["imported"] == 2
    bid2 = upload(uid, (rows + "2024-02-01,Chai,20,expense\n").encode())
    with user_session(uid) as s:
        c = imp.preview(s, uid, bid2, StubFX())["counts"]
    assert (c["duplicate"], c["new"]) == (2, 1)  # third identical chai is new


def test_invalid_rows_reported_and_skipped(uid):
    bid = upload(
        uid,
        b"date,description,amount,type,category\n2024-02-01,Ok,10,expense,\n"
        b"2024-02-30,Bad date,10,expense,\n2024-02-02,Bad cat,10,expense,Nope\n",
    )
    with user_session(uid) as s:
        pv = imp.preview(s, uid, bid, StubFX())
    assert pv["counts"]["invalid"] == 2 and {e["row"] for e in pv["errors"]} == {2, 3}
    assert do_import(uid, bid)["imported"] == 1


def test_undo_removes_exactly_the_batch(uid):
    keep = upload(uid, b"date,description,amount,type\n2024-03-01,Keep me,10,expense\n")
    do_import(uid, keep)
    bid = upload(
        uid,
        b"date,description,amount,type\n2024-03-02,Remove me,10,expense\n2024-03-03,Me too,5,income\n",
    )
    do_import(uid, bid)
    assert count(uid) == 3
    with user_session(uid) as s:
        assert imp.undo(s, uid, bid) == 2
        with pytest.raises(imp.ImportError_):
            imp.undo(s, uid, bid)
    assert count(uid) == 1


def test_bad_uploads_rejected(uid):
    with user_session(uid) as s, pytest.raises(imp.ImportError_):
        imp.create_batch(s, uid, "x.csv", b"")
    with user_session(uid) as s, pytest.raises(imp.ImportError_):
        imp.create_batch(s, uid, "x.xlsx", b"PK\x03\x04garbage")


def test_other_users_batches_are_invisible(uid):
    bid = upload(uid, b"date,description,amount,type\n2024-03-01,x,1,expense\n")
    other = make_user().user_id
    with user_session(other) as s, pytest.raises(imp.ImportError_) as e:
        imp.preview(s, other, bid, StubFX())
    assert e.value.status == 404
