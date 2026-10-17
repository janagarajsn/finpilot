"""Import template: schema JSON, CSV and XLSX (Transactions / Instructions / Categories sheets)
generated from templates/csv/transactions_template.yml + the user's live categories."""

from __future__ import annotations

import csv
import io
from functools import lru_cache
from typing import Any

import yaml

from finpilot.config import get_settings


@lru_cache
def template_def() -> dict[str, Any]:
    path = get_settings().templates_dir / "csv" / "transactions_template.yml"
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def columns() -> list[str]:
    return [c["name"] for c in template_def()["columns"]]


def schema(
    categories: list[dict[str, Any]], currencies: list[str], reporting_currency: str
) -> dict[str, Any]:
    d = template_def()
    return {
        "columns": d["columns"],
        "instructions": d["instructions"],
        "categories": [{"name": c["name"], "flow_type": c["flow_type"]} for c in categories],
        "currencies": currencies,
        "default_currency": reporting_currency,
        "limits": {"max_bytes": 5 * 1024 * 1024, "max_rows": 50_000},
    }


def _examples(reporting_currency: str) -> list[list[str]]:
    ex = template_def()["examples"]
    return [[str(v) for v in row] for row in ex.get(reporting_currency, ex["INR"])]


def build_csv(reporting_currency: str) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(columns())
    w.writerows(_examples(reporting_currency))
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly


def build_xlsx(
    categories: list[dict[str, Any]], currencies: list[str], reporting_currency: str
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    d = template_def()
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    cols = columns()
    ws.append(cols)
    for row in _examples(reporting_currency):
        ws.append(row)
    head = PatternFill("solid", fgColor="2459C9")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    for i, w in enumerate([12, 44, 12, 10, 20, 14, 24, 10, 10], start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "A2"

    cat_ws = wb.create_sheet("Categories")
    cat_ws.append(["name", "flow_type", "description"])
    for c in categories:
        cat_ws.append([c["name"], c["flow_type"], c.get("description", "")])
    cat_ws.column_dimensions["A"].width, cat_ws.column_dimensions["C"].width = 24, 80
    for c in cat_ws[1]:
        c.font = Font(bold=True)

    lists = wb.create_sheet("Lists")  # hidden helper for the currency dropdown
    for i, cur in enumerate(currencies, start=1):
        lists.cell(i, 1, cur)
    lists.sheet_state = "hidden"

    rng = "2:5000"
    idx = {name: ws.cell(1, i + 1).column_letter for i, name in enumerate(cols)}
    dv_type = DataValidation(
        type="list",
        formula1='"expense,income"',
        allow_blank=False,
        showErrorMessage=True,
        error="Choose expense or income",
    )
    dv_cat = DataValidation(
        type="list",
        formula1=f"=Categories!$A$2:$A${len(categories) + 1}",
        allow_blank=True,
        showErrorMessage=True,
        error="Pick a category from the Categories sheet",
    )
    dv_cur = DataValidation(
        type="list", formula1=f"=Lists!$A$1:$A${len(currencies)}", allow_blank=True
    )
    dv_date = DataValidation(
        type="date",
        operator="greaterThan",
        formula1="DATE(1990,1,1)",
        allow_blank=False,
        showErrorMessage=True,
        error="Enter a date (YYYY-MM-DD)",
    )
    dv_amt = DataValidation(
        type="decimal",
        operator="greaterThan",
        formula1="0",
        allow_blank=False,
        showErrorMessage=True,
        error="Amount must be a positive number",
    )
    for dv, col in (
        (dv_type, "type"),
        (dv_cat, "category"),
        (dv_cur, "currency"),
        (dv_date, "date"),
        (dv_amt, "amount"),
    ):
        ws.add_data_validation(dv)
        a, b = rng.split(":")
        dv.add(f"{idx[col]}{a}:{idx[col]}{b}")

    ins = wb.create_sheet("Instructions", 1)
    ins.column_dimensions["A"].width, ins.column_dimensions["B"].width = 14, 100
    ins.append(["FinPilot import template"])
    ins["A1"].font = Font(bold=True, size=14)
    ins.append([])
    for line in d["instructions"]:
        ins.append(["•", line])
    ins.append([])
    ins.append(["Column", "Rules"])
    ins[ins.max_row][0].font = ins[ins.max_row][1].font = Font(bold=True)
    for c in d["columns"]:
        req = "required" if c.get("required") else "optional"
        ins.append([c["name"], f"({req}) {c['description']}"])
    for row in ins.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
