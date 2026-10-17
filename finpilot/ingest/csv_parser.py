"""Tolerant CSV/XLSX reading and row parsing for imports (deterministic, no LLM).

Handles: utf-8(-sig)/cp1252/latin-1, delimiter sniffing, XLSX (first or "Transactions"
sheet), Indian and US grouping (1,25,000.50 / 125,000.50), decimal comma (1.234,56),
currency symbols, parentheses negatives, trailing Dr/Cr, signed amounts or separate
debit/credit columns, and ambiguous dd/mm vs mm/dd dates (reported, user picks).
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import uuid
from collections import Counter
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from functools import lru_cache
from typing import Any

import yaml

from finpilot.config import get_settings

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 50_000
FIELDS = (
    "date",
    "description",
    "amount",
    "debit",
    "credit",
    "type",
    "category",
    "account",
    "notes",
    "currency",
    "fx_rate",
)
DATE_FORMATS = (
    "%Y-%m-%d",
    "%d/%m/%Y",
    "%m/%d/%Y",
    "%d-%m-%Y",
    "%m-%d-%Y",
    "%d.%m.%Y",
    "%Y/%m/%d",
    "%d-%b-%Y",
    "%d %b %Y",
    "%b %d, %Y",
    "%d-%b-%y",
    "%d/%m/%y",
    "%m/%d/%y",
    "%Y%m%d",
)
AMBIGUOUS_PAIRS = {("%d/%m/%Y", "%m/%d/%Y"), ("%d-%m-%Y", "%m-%d-%Y"), ("%d/%m/%y", "%m/%d/%y")}


class ImportFileError(ValueError):
    pass


@lru_cache
def _mapping_cfg() -> dict[str, Any]:
    return yaml.safe_load(
        (get_settings().config_dir / "import_mapping.yml").read_text(encoding="utf-8")
    )


def _norm_header(h: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[_\-]+", " ", str(h or ""))).strip().lower()


# ---------------------------------------------------------------------------- reading
@dataclass
class Table:
    headers: list[str]
    rows: list[list[str]]
    kind: str  # csv | xlsx
    encoding: str | None = None
    delimiter: str | None = None


def _decode(content: bytes) -> tuple[str, str]:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return content.decode(enc), enc
        except UnicodeDecodeError:
            continue
    raise ImportFileError("Could not read the file's text encoding.")


def sniff_delimiter(sample: str) -> str:
    """Pick the delimiter that splits the most lines into the SAME number (>1) of fields.
    More robust than csv.Sniffer with bank preambles and decimal-comma amounts."""
    lines = [ln for ln in sample.splitlines() if ln.strip()][:60]
    best, best_score = ",", -1
    for d in (",", ";", "\t", "|"):
        counts = Counter(len(r) for r in csv.reader(lines, delimiter=d))
        width, n = max(
            ((w, c) for w, c in counts.items() if w > 1), key=lambda x: (x[1], x[0]), default=(1, 0)
        )
        score = n * 1000 + width
        if score > best_score:
            best, best_score = d, score
    return best


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float):
        return repr(v) if not v.is_integer() else str(int(v))
    return str(v).strip()


def read_table(content: bytes, filename: str = "") -> Table:
    if len(content) > MAX_BYTES:
        raise ImportFileError("File is larger than 5 MB.")
    if not content.strip():
        raise ImportFileError("The file is empty.")
    is_xlsx = content[:2] == b"PK" or filename.lower().endswith((".xlsx", ".xlsm"))
    if is_xlsx:
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as e:
            raise ImportFileError("Could not open the Excel file.") from e
        ws = wb["Transactions"] if "Transactions" in wb.sheetnames else wb.worksheets[0]
        raw_rows = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        table = Table(headers=[], rows=[], kind="xlsx")
    else:
        text_, enc = _decode(content)
        delim = sniff_delimiter(text_[:50_000])
        raw_rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text_), delimiter=delim)]
        table = Table(headers=[], rows=[], kind="csv", encoding=enc, delimiter=delim)
    raw_rows = [r for r in raw_rows if any(c for c in r)]
    if not raw_rows:
        raise ImportFileError("No rows found.")
    # Header = first row with at least two non-empty, mostly non-numeric cells (skips bank preambles).
    hdr_idx = 0
    for i, r in enumerate(raw_rows[:15]):
        labels = [c for c in r if c and not re.fullmatch(r"[\d.,\-/ ]+", c)]
        if len(labels) >= 2 and any(_guess_field(c) for c in r):
            hdr_idx = i
            break
    headers = [h or f"column_{i + 1}" for i, h in enumerate(raw_rows[hdr_idx])]
    width = len(headers)
    rows = [(r + [""] * width)[:width] for r in raw_rows[hdr_idx + 1 :]]
    if len(rows) > MAX_ROWS:
        raise ImportFileError(f"Too many rows ({len(rows):,}). The limit is {MAX_ROWS:,}.")
    table.headers, table.rows = headers, rows
    return table


def _guess_field(header: str) -> str | None:
    h = _norm_header(header)
    for f, syns in _mapping_cfg()["fields"].items():
        if h == f or h in syns:
            return str(f)
    return None


def guess_mapping(headers: list[str]) -> dict[str, int | None]:
    mapping: dict[str, int | None] = dict.fromkeys(FIELDS)
    for i, h in enumerate(headers):
        f = _guess_field(h)
        if f and mapping.get(f) is None:
            mapping[f] = i
    if mapping["amount"] is None and mapping["debit"] is None and mapping["credit"] is None:
        for i, h in enumerate(headers):  # loose fallback: any header containing "amount"
            if "amount" in _norm_header(h):
                mapping["amount"] = i
                break
    return mapping


# ---------------------------------------------------------------------------- values
_CUR_TOKENS = re.compile(r"(₹|rs\.?|inr|usd|us\$|\$)", re.IGNORECASE)


def parse_amount(raw: str, decimal_comma: bool = False) -> tuple[Decimal | None, int]:
    """Returns (absolute or signed amount, sign_hint) — sign_hint is -1 (Dr), +1 (Cr) or 0."""
    s = (raw or "").strip()
    if not s:
        return None, 0
    hint = 0
    m = re.search(r"\b(dr|cr)\.?$", s, re.IGNORECASE)
    if m:
        hint = -1 if m.group(1).lower() == "dr" else 1
        s = s[: m.start()].strip()
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1]
    s = _CUR_TOKENS.sub("", s).replace(" ", "").replace(" ", "")
    if s.startswith("-"):
        neg, s = not neg, s[1:]
    elif s.startswith("+"):
        s = s[1:]
    if s.endswith("-"):  # trailing minus (some bank exports)
        neg, s = not neg, s[:-1]
    s = s.replace(".", "").replace(",", ".") if decimal_comma else s.replace(",", "")
    if not re.fullmatch(r"\d+(\.\d+)?|\.\d+", s):
        return None, hint
    try:
        val = Decimal(s)
    except InvalidOperation:
        return None, hint
    return (-val if neg else val), hint


def looks_decimal_comma(values: list[str]) -> bool:
    """True if amounts look like 1.234,56 / 12,50 (European)."""
    eu = sum(bool(re.fullmatch(r"-?[\d.]*\d,\d{1,2}", v.strip())) for v in values if v.strip())
    us = sum(bool(re.fullmatch(r"-?[\d,]*\d\.\d{1,2}", v.strip())) for v in values if v.strip())
    return eu > us and eu > 0


def detect_date_formats(values: list[str]) -> list[str]:
    vals = [v.strip() for v in values if v and v.strip()][:2000]
    if not vals:
        return []
    out = []
    for fmt in DATE_FORMATS:
        try:
            for v in vals:
                dt.datetime.strptime(v, fmt)
            out.append(fmt)
        except ValueError:
            continue
    return out


def is_ambiguous(formats: list[str]) -> bool:
    return any(a in formats and b in formats for a, b in AMBIGUOUS_PAIRS)


# ---------------------------------------------------------------------------- rows
@dataclass
class ParseContext:
    category_index: dict[str, uuid.UUID]  # lower(name) → id
    account_index: dict[str, tuple[uuid.UUID, str]]  # lower(name) → (id, currency)
    supported_currencies: set[str]
    reporting_currency: str
    today: dt.date


@dataclass
class ParsedRow:
    row_no: int  # 1-based data row number (header excluded)
    raw: list[str]
    data: dict[str, Any] | None = None
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _get(row: list[str], idx: int | None) -> str:
    return row[idx].strip() if idx is not None and 0 <= idx < len(row) else ""


def parse_rows(
    table: Table, mapping: dict[str, int | None], options: dict[str, Any], ctx: ParseContext
) -> list[ParsedRow]:
    date_fmt: str | None = options.get("date_format")
    decimal_comma = bool(options.get("decimal_comma"))
    sign_convention = options.get("amount_sign", "negative_is_expense")
    type_vals = _mapping_cfg()["type_values"]
    out: list[ParsedRow] = []
    for i, row in enumerate(table.rows, start=1):
        pr = ParsedRow(i, row)
        # date
        d_raw = _get(row, mapping.get("date"))
        date: dt.date | None = None
        if not d_raw:
            pr.errors.append("Date is missing.")
        else:
            fmts = [date_fmt] if date_fmt else list(DATE_FORMATS)
            for fmt in fmts:
                try:
                    date = dt.datetime.strptime(d_raw, fmt).date()
                    break
                except ValueError:
                    continue
            if date is None:
                pr.errors.append(f"Unrecognized date “{d_raw}”.")
            elif date > ctx.today + dt.timedelta(days=366):
                pr.errors.append("Date is more than a year in the future.")
            elif date < dt.date(1990, 1, 1):
                pr.errors.append("Date is before 1990.")
            elif date > ctx.today:
                pr.warnings.append("Future-dated transaction.")
        # description
        desc = re.sub(r"\s+", " ", _get(row, mapping.get("description")))
        if not desc:
            pr.errors.append("Description is missing.")
        elif len(desc) > 500:
            pr.errors.append("Description is longer than 500 characters.")
        # amount + direction
        kind: str | None = None
        amount: Decimal | None = None
        t_raw = _get(row, mapping.get("type")).lower()
        both_filled = False
        if mapping.get("debit") is not None or mapping.get("credit") is not None:
            dr, _ = parse_amount(_get(row, mapping.get("debit")), decimal_comma)
            cr, _ = parse_amount(_get(row, mapping.get("credit")), decimal_comma)
            if dr and dr != 0 and cr and cr != 0:
                both_filled = True
                pr.errors.append("Both debit and credit are filled.")
            elif dr and dr != 0:
                kind, amount = "expense", abs(dr)
            elif cr and cr != 0:
                kind, amount = "income", abs(cr)
            elif mapping.get("amount") is None:
                pr.errors.append("No debit or credit amount.")
        if kind is None and mapping.get("amount") is not None and not both_filled:
            a_raw = _get(row, mapping.get("amount"))
            val, hint = parse_amount(a_raw, decimal_comma)
            if val is None:
                pr.errors.append(
                    f"Amount “{a_raw}” is not a number." if a_raw else "Amount is missing."
                )
            elif val == 0:
                pr.errors.append("Amount is zero.")
            else:
                amount = abs(val)
                if t_raw:
                    if t_raw in type_vals["expense"]:
                        kind = "expense"
                    elif t_raw in type_vals["income"]:
                        kind = "income"
                    else:
                        pr.errors.append(f"Type “{t_raw}” must be expense or income.")
                    if val < 0 and kind == "income":
                        pr.warnings.append("Negative amount with type income: treated as income.")
                elif hint:
                    kind = "expense" if hint < 0 else "income"
                else:
                    neg_is_exp = sign_convention == "negative_is_expense"
                    kind = (
                        ("expense" if val < 0 else "income")
                        if neg_is_exp
                        else ("income" if val < 0 else "expense")
                    )
        if amount is not None and amount >= Decimal("1e14"):
            pr.errors.append("Amount is too large.")
        # account → default currency
        account_id, account_currency = None, None
        acct_raw = _get(row, mapping.get("account"))
        if acct_raw:
            hit = ctx.account_index.get(acct_raw.lower())
            if hit:
                account_id, account_currency = hit
            else:
                pr.warnings.append(f"Unknown account “{acct_raw}”: imported without an account.")
        # currency
        cur_raw = _get(row, mapping.get("currency")).upper()
        currency = cur_raw or account_currency or ctx.reporting_currency
        if currency not in ctx.supported_currencies:
            pr.errors.append(f"Currency “{currency}” is not supported.")
        # fx_rate override
        fx_rate: Decimal | None = None
        fx_raw = _get(row, mapping.get("fx_rate"))
        if fx_raw:
            v, _ = parse_amount(fx_raw)
            if v is None or v <= 0:
                pr.errors.append(f"fx_rate “{fx_raw}” must be a positive number.")
            elif currency == "INR":
                pr.warnings.append("fx_rate ignored for INR rows.")
            else:
                fx_rate = v
        # category
        category_id = None
        cat_raw = _get(row, mapping.get("category"))
        if cat_raw:
            category_id = ctx.category_index.get(re.sub(r"\s+", " ", cat_raw).lower())
            if category_id is None:
                pr.errors.append(
                    f"Category “{cat_raw}” does not exist. Create it first or leave the cell blank."
                )
        notes = _get(row, mapping.get("notes"))[:1000] or None
        if not pr.errors and kind and amount is not None and date is not None:
            pr.data = {
                "kind": kind,
                "amount": amount,
                "date": date,
                "description": desc,
                "currency": currency,
                "category_id": category_id,
                "account_id": account_id,
                "notes": notes,
                "fx_rate": fx_rate,
            }
        elif not pr.errors:
            pr.errors.append("Could not determine amount and direction.")
        out.append(pr)
    return out


def validate_mapping(mapping: dict[str, Any], width: int) -> dict[str, int | None]:
    clean: dict[str, int | None] = dict.fromkeys(FIELDS)
    for k, v in (mapping or {}).items():
        if k not in FIELDS:
            raise ImportFileError(f"Unknown field “{k}”.")
        if v in (None, "", -1):
            continue
        iv = int(v)
        if not 0 <= iv < width:
            raise ImportFileError(f"Column for “{k}” is out of range.")
        clean[k] = iv
    if clean["date"] is None or clean["description"] is None:
        raise ImportFileError("Map at least the date and description columns.")
    if clean["amount"] is None and clean["debit"] is None and clean["credit"] is None:
        raise ImportFileError("Map an amount column (or debit/credit columns).")
    used = [v for v in clean.values() if v is not None]
    if len(used) != len(set(used)):
        raise ImportFileError("Each column can be mapped to only one field.")
    return clean
